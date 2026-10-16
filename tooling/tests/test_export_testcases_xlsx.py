"""Exporter tests for the Sprint Number display and workbook shape."""

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "export_testcases_xlsx.py"
SPEC = importlib.util.spec_from_file_location("export_testcases_xlsx", SCRIPT)
exporter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(exporter)


class ExportTestCasesTests(unittest.TestCase):
    def _parse_sprint(self, metadata: list[str]) -> object:
        block = "\n".join([
            "### TC_E01S01_01 — Example case",
            *metadata,
            "#### Pre-conditions",
            "_None recorded._",
        ])
        return exporter._parse_case_block(block, "E01-S01")["sprint"]

    def test_empty_sprint_displays_unassigned(self):
        self.assertEqual(self._parse_sprint(["- **Sprint**:"]), "Unassigned")

    def test_whitespace_only_sprint_displays_unassigned(self):
        self.assertEqual(self._parse_sprint(["- **Sprint**:   \t"]), "Unassigned")

    def test_missing_sprint_displays_unassigned(self):
        self.assertEqual(self._parse_sprint([]), "Unassigned")

    def test_numeric_sprint_stays_numeric(self):
        self.assertEqual(self._parse_sprint(["- **Sprint**: 3"]), 3)

    def test_fractional_numeric_sprint_stays_numeric(self):
        self.assertEqual(self._parse_sprint(["- **Sprint**: 3.5"]), 3.5)

    def test_nonnumeric_sprint_keeps_its_text(self):
        value = "Follow-up; planning pending"
        self.assertEqual(self._parse_sprint([f"- **Sprint**: {value}"]), value)

    def test_export_preserves_catalogue_shape_ids_and_non_sprint_values(self):
        case_files = sorted(exporter.CASES_DIR.glob("*.md"))
        rows = [case for path in case_files for case in exporter._parse_epic_file(path)]
        self.assertTrue(rows)
        expected_unassigned_by_story = {}
        for case in rows:
            if case["sprint"] == exporter.UNASSIGNED_SPRINT_DISPLAY:
                story_id = case["story_id"]
                expected_unassigned_by_story[story_id] = (
                    expected_unassigned_by_story.get(story_id, 0) + 1
                )

        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "test-cases.xlsx"
            exporter._write_workbook(rows, destination)
            workbook = load_workbook(destination, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["TEST CASES"])
                sheet = workbook["TEST CASES"]
                self.assertEqual(sheet.max_column, 10)
                self.assertEqual(sheet.max_row, len(rows) + 1)
                self.assertEqual(sheet.column_dimensions["A"].width, 14)
                self.assertEqual(tuple(next(sheet.iter_rows(values_only=True))), exporter.COLUMNS)

                exported = list(sheet.iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(exported), len(rows))
                self.assertEqual(
                    [row[4] for row in exported],
                    [case["case_id"] for case in rows],
                )
                self.assertEqual(
                    sum(row[0] == exporter.UNASSIGNED_SPRINT_DISPLAY for row in exported),
                    sum(expected_unassigned_by_story.values()),
                )
                unassigned_by_story = {}
                for row in exported:
                    if row[0] == exporter.UNASSIGNED_SPRINT_DISPLAY:
                        unassigned_by_story[row[1]] = unassigned_by_story.get(row[1], 0) + 1
                self.assertEqual(
                    unassigned_by_story,
                    expected_unassigned_by_story,
                )

                expected_non_sprint = [
                    tuple(
                        None if value == "" else value
                        for value in (
                            case["story_id"], case["ac_reference"], case["step_type"],
                            case["case_id"], case["scenario"], case["preconditions"],
                            case["steps"], case["test_data"], case["expected"],
                        )
                    )
                    for case in rows
                ]
                self.assertEqual([tuple(row[1:]) for row in exported], expected_non_sprint)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
