"""One-shot: convert docs/testing/PROJECT TEST CASES.xlsx to per-epic
Markdown files under docs/testing/cases/. Same pattern as
scripts/one_shot/backlog_xlsx_to_md.py. Requires openpyxl."""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCE_XLSX = REPO_ROOT / "docs" / "testing" / "PROJECT TEST CASES.xlsx"
CASES_DIR = REPO_ROOT / "docs" / "testing" / "cases"


def _epic_id(story_id: str) -> str:
    match = re.match(r"^(E\d{2})", story_id.strip())
    return match.group(1) if match else "EXX"


def _slug(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower() or "epic"


def _read_test_cases(worksheet) -> list[dict]:
    header = [str(cell.value).strip() if cell.value else "" for cell in worksheet[1]]
    field_index = {name: index for index, name in enumerate(header)}
    field_map = {
        "sprint":       "Sprint Number",
        "story_id":     "Story ID",
        "ac_reference": "AC Reference",
        "step_type":    "Step(1-5)",
        "case_id":      "Test Case ID",
        "scenario":     "Test Case Scenerio",  # workbook uses this exact spelling
        "preconditions":"Pre Conditions",
        "steps":        "Test Steps",
        "test_data":    "Test Data",
        "expected":     "Expected Result",
    }
    rows: list[dict] = []
    for raw in worksheet.iter_rows(min_row=2, values_only=True):
        if not raw or all(cell in (None, "") for cell in raw):
            continue
        record = {}
        for key, header_name in field_map.items():
            index = field_index.get(header_name)
            record[key] = raw[index] if index is not None and index < len(raw) else None
        if not record.get("case_id"):
            continue
        rows.append(record)
    return rows


def _render_case(record: dict) -> str:
    case_id = str(record["case_id"]).strip()
    scenario = str(record.get("scenario") or "").strip()
    sprint_raw = record.get("sprint")
    sprint = "" if sprint_raw in (None, "") else (
        str(int(sprint_raw)) if isinstance(sprint_raw, float) and sprint_raw.is_integer() else str(sprint_raw)
    )
    parts: list[str] = []
    parts.append(f"### {case_id} — {scenario or 'Test case'}")
    parts.append("")
    parts.append(f"- **Sprint**: {sprint}")
    parts.append(f"- **AC reference**: {(record.get('ac_reference') or '').strip()}")
    parts.append(f"- **Type**: {(record.get('step_type') or '').strip()}")
    parts.append("")
    parts.append("#### Pre-conditions")
    parts.append("")
    parts.append((record.get("preconditions") or "").strip() or "_None recorded._")
    parts.append("")
    parts.append("#### Test steps")
    parts.append("")
    parts.append((record.get("steps") or "").strip() or "_None recorded._")
    parts.append("")
    parts.append("#### Test data")
    parts.append("")
    parts.append((record.get("test_data") or "").strip() or "_None recorded._")
    parts.append("")
    parts.append("#### Expected result")
    parts.append("")
    parts.append((record.get("expected") or "").strip() or "_None recorded._")
    parts.append("")
    return "\n".join(parts)


def _write_epic_file(target_dir: Path, epic_id: str, rows: Iterable[dict]) -> None:
    target_path = target_dir / f"{epic_id}.md"
    parts: list[str] = []
    parts.append(f"# {epic_id} test cases")
    parts.append("")
    ordered = sorted(rows, key=lambda r: (str(r.get("story_id") or ""), str(r.get("case_id") or "")))
    current_story = None
    for row in ordered:
        story_id = str(row.get("story_id") or "").strip() or "Unassigned"
        if story_id != current_story:
            parts.append(f"## {story_id}")
            parts.append("")
            current_story = story_id
        parts.append(_render_case(row))
    target_path.write_text("\n".join(parts), encoding="utf-8")


def main() -> int:
    if not SOURCE_XLSX.is_file():
        print(f"ERROR: source workbook not found: {SOURCE_XLSX}", file=sys.stderr)
        return 1
    workbook = openpyxl.load_workbook(str(SOURCE_XLSX), data_only=True)
    if "TEST CASES" not in workbook.sheetnames:
        print("ERROR: sheet 'TEST CASES' missing", file=sys.stderr)
        return 1
    rows = _read_test_cases(workbook["TEST CASES"])
    CASES_DIR.mkdir(parents=True, exist_ok=True)
    grouped: dict[str, list[dict]] = {}
    for row in rows:
        epic = _epic_id(str(row.get("story_id") or ""))
        grouped.setdefault(epic, []).append(row)
    for epic_id, epic_rows in sorted(grouped.items()):
        _write_epic_file(CASES_DIR, epic_id, epic_rows)
    print(f"Wrote {len(grouped)} epic file(s), {len(rows)} test cases total")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
